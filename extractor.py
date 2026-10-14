"""Module 1: Resume upload & text extraction"""
try:
    import pymupdf as fitz  # PyMuPDF (new name)
except ImportError:
    import fitz  # PyMuPDF (old name)
import docx  # python-docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def _table_text(table, parts, seen):
    """Add the text of a table (and tables inside it). Many resume templates keep skills in tables."""
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen:  # merged cells come back more than once
                continue
            seen.add(cell._tc)
            for para in cell.paragraphs:
                parts.append(para.text)
            for inner in cell.tables:
                _table_text(inner, parts, seen)


def _docx_text(path):
    document = docx.Document(path)
    parts, seen = [], set()
    for child in document.element.body.iterchildren():  # keep the order of the document
        if child.tag == qn("w:p"):
            parts.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            _table_text(Table(child, document), parts, seen)
    return "\n".join(parts)


def extract_text(path):
    """Return the text of a PDF/DOCX file, or "" if it cannot be read (scanned, empty, damaged, password protected)."""
    p = path.lower()
    try:
        if p.endswith(".pdf"):
            with fitz.open(path) as pdf:
                if pdf.needs_pass:
                    return ""
                return "\n".join(page.get_text() for page in pdf)
        if p.endswith(".docx"):
            return _docx_text(path)
    except Exception:
        return ""
    return ""
